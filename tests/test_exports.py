from io import BytesIO
import unittest
from openpyxl import load_workbook
from exports import excel_bytes, course_and_year, group_sort_key, surname_first


class ExportTests(unittest.TestCase):
    def test_reference_layout_colors_and_formula_safety(self):
        roster = [{"email": "a", "name": "=1+1", "degree": 2}]
        groups = [
            {
                "id": 1,
                "period": 9,
                "degree": 2,
                "name": "Math",
                "professor": "Prof Math",
                "room": "A",
                "kind": "remediation",
            },
            {
                "id": 2,
                "period": 10,
                "degree": 2,
                "name": "Étude",
                "professor": "Prof Étude",
                "room": "B",
                "kind": "depassement",
            },
        ]
        assignments = [
            {"email": "a", "period": 9, "activity_id": 1},
            {"email": "a", "period": 10, "activity_id": 2},
        ]
        wb = load_workbook(
            BytesIO(
                excel_bytes(
                    {"title": "Test", "event_date": "2026-09-15"},
                    roster,
                    groups,
                    assignments,
                )
            )
        )
        self.assertEqual(wb.sheetnames, ["D1", "D2", "D3"])
        ws = wb["D2"]
        self.assertEqual(ws["B1"].value, "\nMath\nProf Math\nLocal A")
        self.assertEqual(ws["B2"].data_type, "s")
        self.assertEqual(ws["C19"].value, "\nÉtude\nProf Étude\nLocal B")
        self.assertEqual(ws["C20"].value, "=1+1")
        self.assertEqual(ws["A1"].value, "P9")
        self.assertEqual(ws["A19"].value, "P10")
        self.assertEqual(ws["A2"].value, 1)
        self.assertEqual(ws["A20"].value, 1)
        self.assertFalse(ws.merged_cells.ranges)
        self.assertIsNone(ws["B2"].fill.patternType)
        self.assertIsNone(ws["B3"].fill.patternType)
        self.assertTrue(ws.sheet_view.showGridLines)
        self.assertTrue(ws.print_options.gridLines)
        self.assertEqual(ws.page_setup.orientation, "landscape")
        for sheet in wb:
            for item in ("oddHeader", "evenHeader", "firstHeader", "oddFooter", "evenFooter", "firstFooter"):
                self.assertFalse(getattr(sheet, item))
        self.assertEqual(ws.column_dimensions["B"].width, 27)
        self.assertEqual(ws.column_dimensions["C"].width, 27)

    def test_unregistered_per_period_and_all_empty_activities_exported(self):
        roster = [
            {"email": "a", "name": "Alice Zed", "degree": 2},
            {"email": "b", "name": "Bob Alpha", "degree": 2},
            {"email": "c", "name": "Cleo Autre", "degree": 1},
        ]
        groups = [
            dict(id=1, period=9, degree=2, name="Math", professor="Prof", room="A"),
            dict(id=2, period=10, degree=2, name="Math", professor="Prof", room="B"),
            dict(id=3, period=9, degree=2, name="Vide", professor="Prof Vide", room="C"),
        ]
        assignments = [{"email": "a", "period": 9, "activity_id": 1}]
        wb = load_workbook(BytesIO(excel_bytes(
            {"title": "Test", "event_date": "2026-09-15"},
            roster, groups, assignments, skip_unregistered=True,
        )))
        ws = wb["D2"]
        self.assertEqual(ws["B2"].value, "Zed Alice")
        self.assertEqual(ws["C1"].value, "\nVide\nProf Vide\nLocal C")
        self.assertIsNone(ws["C2"].value)
        self.assertEqual(ws["C19"].value, "Non proposée en P10")
        self.assertEqual(ws["B19"].value, "\nMath\nProf\nLocal B")
        self.assertEqual(ws["D1"].value, "Non inscrits")
        self.assertEqual(ws["D19"].value, "Non inscrits")
        self.assertEqual(ws["D2"].value, "Alpha Bob")
        self.assertIsNone(ws["D3"].value)
        self.assertEqual(ws["D20"].value, "Alpha Bob")
        self.assertEqual(ws["D21"].value, "Zed Alice")
        for column in ("B", "C", "D"):
            self.assertEqual(ws.column_dimensions[column].width, 27)
        self.assertIn("$D$31", str(ws.print_area))
        self.assertEqual(wb["D1"]["B2"].value, "Autre Cleo")
        self.assertEqual(wb["D1"]["B20"].value, "Autre Cleo")
        self.assertEqual(wb["D3"]["B1"].value, "Non inscrits")
        self.assertIsNone(wb["D3"]["B2"].value)

    def test_unregistered_column_kept_when_everyone_is_registered(self):
        groups = [
            dict(id=1, period=9, degree=2, name="Vide", professor="Prof", room="A"),
            dict(id=2, period=10, degree=2, name="Vide", professor="Prof", room="B"),
        ]
        for enabled in (False, True):
            with self.subTest(enabled=enabled):
                wb = load_workbook(BytesIO(excel_bytes(
                    {"title": "Test", "event_date": "2026-09-15"},
                    [], groups, [], skip_unregistered=enabled,
                )))
                ws = wb["D2"]
                self.assertEqual(ws["B1"].value, "\nVide\nProf\nLocal A")
                self.assertEqual(ws["B19"].value, "\nVide\nProf\nLocal B")
                self.assertIsNone(ws["B2"].value)
                self.assertEqual(ws.max_column, 3 if enabled else 2)
                if enabled:
                    self.assertEqual(ws["C1"].value, "Non inscrits")
                    self.assertIsNone(ws["C2"].value)

    def test_unregistered_not_truncated_without_activities_and_formula_safe(self):
        roster = [
            {"email": str(i), "name": f"Eleve {i:02}", "degree": 1}
            for i in range(30)
        ] + [{"email": "formula", "name": "=1+1", "degree": 1}]
        wb = load_workbook(BytesIO(excel_bytes(
            {"title": "Test", "event_date": "2026-09-15"},
            roster, [], [], skip_unregistered=True,
        )))
        ws = wb["D1"]
        self.assertEqual(ws["B1"].value, "Non inscrits")
        self.assertEqual(ws["B31"].value, "29 Eleve")
        self.assertEqual(ws["B32"].value, "=1+1")
        self.assertEqual(ws["B32"].data_type, "s")
        self.assertEqual(ws["A38"].value, "P10")
        self.assertEqual(ws["B69"].value, "=1+1")
        self.assertIn("$B$69", str(ws.print_area))

    def test_activity_named_non_inscrits_has_its_own_column(self):
        group = dict(id=1, period=9, degree=1, name="Non inscrits", professor="Prof", room="A")
        wb = load_workbook(BytesIO(excel_bytes(
            {"title": "Test", "event_date": "2026-09-15"},
            [{"email": "a", "name": "Alice Dupont", "degree": 1}], [group],
            [{"email": "a", "period": 9, "activity_id": 1}], skip_unregistered=True,
        )))
        self.assertEqual(wb["D1"]["B2"].value, "Dupont Alice")
        self.assertEqual(wb["D1"]["C1"].value, "Non inscrits")
        self.assertIsNone(wb["D1"]["C2"].value)
        self.assertEqual(wb["D1"]["C20"].value, "Dupont Alice")

    def test_unlimited_group_never_truncated_and_next_block_moves(self):
        roster = [
            {"email": str(i), "name": f"Élève {i:02}", "degree": 1} for i in range(30)
        ]
        groups = [
            {
                "id": i,
                "degree": 1,
                "period": 8 + i,
                "name": f"Activité {i}",
                "professor": "Prof",
                "room": "A",
            }
            for i in (1, 2)
        ]
        assignments = [
            {"email": str(i), "activity_id": a, "period": 8 + a}
            for i in range(30)
            for a in (1, 2)
        ]
        wb = load_workbook(
            BytesIO(
                excel_bytes(
                    {"title": "Test", "event_date": "2026-09-15"},
                    roster,
                    groups,
                    assignments,
                )
            )
        )
        self.assertEqual(wb["D1"]["B31"].value, "29 Élève")
        self.assertEqual(wb["D1"]["A37"].value, "P10")
        self.assertEqual(wb["D1"]["C67"].value, "29 Élève")
        self.assertEqual(wb["D1"]["A67"].value, 30)

    def test_year_extraction_and_group_order(self):
        for name, expected in (
            ("Math 2ème", ("Math", "2ème", 2)),
            ("Français 1", ("Français", "1ère", 1)),
            ("Sciences 3e", ("Sciences", "3ème", 3)),
            ("Math 4ième", ("Math", "4ème", 4)),
            ("Math 5ème", ("Math", "5ème", 5)),
            ("Math 6ème", ("Math", "6ème", 6)),
            ("Robotique", ("Robotique", "", None)),
        ):
            with self.subTest(name=name):
                self.assertEqual(course_and_year(name), expected)
        names = [
            ("ÉTUDE accompagnée", "depassement"),
            ("Math 2ème", "remediation"),
            ("Atelier", "depassement"),
            ("Sciences 1ère", "remediation"),
            ("Math", "remediation"),
            ("Zoologie", "depassement"),
        ]
        groups = [{"name": n, "kind": k, "period": 9} for n, k in names]
        self.assertEqual(
            [g["name"] for g in sorted(groups, key=group_sort_key)],
            ["Sciences 1ère", "Math 2ème", "Math", "Atelier", "Zoologie", "ÉTUDE accompagnée"],
        )

    def test_surname_first_preserves_family_name(self):
        self.assertEqual(surname_first("Alice de la Fontaine"), "de la Fontaine Alice")
        self.assertEqual(surname_first("Jean-Pierre Dupont"), "Dupont Jean-Pierre")
        self.assertEqual(surname_first("Alice"), "Alice")

    def test_incomplete_excel_refused(self):
        with self.assertRaises(ValueError):
            excel_bytes(
                {"title": "Test", "event_date": "2026-09-15"},
                [{"email": "a", "name": "A", "degree": 2}],
                [],
                [],
            )
