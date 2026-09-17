from datetime import date

"""Export des groupes en colonnes, selon le modèle Focus Time fourni."""

from io import BytesIO
import math
import re
import unicodedata

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

from activity_import import activity_name

COLORS = ("83CCEB", "F6C6AD", "84E291", "9EDCF3", "E59DDD", "B7E5A5")

YEAR = re.compile(r"\b([1-6])(?:\s*(?:ères?|eres?|ers?|ièmes?|iemes?|èmes?|emes?|es?))?\b", re.IGNORECASE)


def course_and_year(name):
    match = YEAR.search(name)
    if not match:
        return name.strip(), "", None
    year = int(match.group(1))
    course = (name[:match.start()] + name[match.end():]).strip(" -–—()")
    return " ".join(course.split()), "1ère" if year == 1 else f"{year}ème", year


def group_sort_key(group):
    course, _, year = course_and_year(group["name"])
    normalized = "".join(
        c for c in unicodedata.normalize("NFD", group["name"].casefold())
        if not unicodedata.combining(c)
    )
    category = 0 if group.get("kind", "remediation") == "remediation" else 1
    if category and re.search(r"\betude\b", normalized):
        category = 2
    return category, (year or 99) if category == 0 else 0, course.casefold(), activity_name(group["name"]), group["period"]


def surname_first(name):
    # The roster stores a given name followed by the full family name.
    parts = name.strip().split(maxsplit=1)
    return f"{parts[1]} {parts[0]}" if len(parts) == 2 else name.strip()


def excel_bytes(session, roster, activities, assignments):
    activity_map = {a["id"]: a for a in activities}
    roster_map = {s["email"]: s for s in roster}
    assigned = {(a["email"], a["period"]) for a in assignments}
    if any((s["email"], p) not in assigned for s in roster for p in (9, 10)):
        raise ValueError("Le planning est incomplet : génération de l'Excel refusée.")
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.title = "{} — {}".format(
        session["title"], date.fromisoformat(session["event_date"]).strftime("%d/%m/%Y")
    )
    for degree in (1, 2, 3):
        ws = wb.create_sheet(f"D{degree}")
        ws.sheet_view.showGridLines = True
        ws.print_options.gridLines = True
        ws.column_dimensions["A"].width = 10.66
        groups = sorted(
            [a for a in activities if a["degree"] == degree],
            key=group_sort_key,
        )
        names = list(dict.fromkeys(activity_name(a["name"]) for a in groups))
        if not groups:
            ws["A1"] = f"D{degree} — Aucune activité"
            ws["A1"].font = Font(name="Arial", size=11)
            ws.merge_cells("A1:D1")
            ws.print_area = "A1:D1"
        header_row = 1
        for period in (9, 10) if groups else ():
            members = {}
            by_name = {
                activity_name(a["name"]): a for a in groups if a["period"] == period
            }
            for name, group in by_name.items():
                members[name] = sorted(
                    [
                        surname_first(roster_map[r["email"]]["name"])
                        for r in assignments
                        if r["activity_id"] == group["id"]
                    ],
                    key=str.casefold,
                )
            row_count = max(12, max((len(v) for v in members.values()), default=0))
            first_row, last_row = header_row + 1, header_row + row_count
            label = ws.cell(header_row, 1, f"P{period}")
            label.font = Font(name="Arial", size=28, bold=True)
            label.alignment = Alignment(horizontal="center", vertical="center")
            for offset in range(row_count):
                number = ws.cell(first_row + offset, 1, offset + 1)
                number.font = Font(name="Arial", size=11)
                number.alignment = Alignment(horizontal="center", vertical="center")
            header_height = 76
            for column, name in enumerate(names, 2):
                ws.column_dimensions[get_column_letter(column)].width = 30.66
                group = by_name.get(name)
                header = ws.cell(header_row, column)
                if group:
                    course, year_label, _ = course_and_year(group["name"])
                    title = "{}\n{}\n{}\nLocal {}".format(
                        year_label, course, group["professor"], group["room"]
                    )
                    header.value = title
                    header.data_type = "s"
                    header.fill = PatternFill(
                        "solid", fgColor=COLORS[(column - 2) % len(COLORS)]
                    )
                    lines = sum(
                        max(1, math.ceil(len(part) / 30)) for part in title.split("\n")
                    )
                    header_height = max(header_height, lines * 14 + 16)
                else:
                    header.value = "Non proposée en P{}".format(period)
                    header.fill = PatternFill("solid", fgColor="ECECEC")
                header.font = Font(name="Arial", size=11, bold=True)
                header.alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True
                )
                for offset in range(row_count):
                    cell = ws.cell(first_row + offset, column)
                    cell.font = Font(name="Arial", size=11)
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
                    if offset < len(members.get(name, [])):
                        cell.value = members[name][offset]
                        cell.data_type = "s"
                    if offset == row_count - 1:
                        cell.border = Border(bottom=Side(style="thin", color="808080"))
            ws.row_dimensions[header_row].height = header_height
            for row in range(first_row, last_row + 1):
                values = [
                    str(ws.cell(row, col).value or "")
                    for col in range(2, len(names) + 2)
                ]
                lines = max(
                    (max(1, math.ceil(len(value) / 32)) for value in values), default=1
                )
                ws.row_dimensions[row].height = max(24.9, lines * 14 + 6)
            header_row = last_row + 6
            ws.print_area = "A1:{}{}".format(
                get_column_letter(len(names) + 1), last_row
            )
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_margins = PageMargins(left=0.25, right=0.25, top=0.5, bottom=0.5)
        ws.print_options.horizontalCentered = True
        ws.oddHeader.center.text = "{} — {} — D{}".format(
            session["title"],
            date.fromisoformat(session["event_date"]).strftime("%d/%m/%Y"),
            degree,
        )
        ws.oddFooter.right.text = "Page &P / &N"
    output = BytesIO()
    wb.save(output)
    return output.getvalue()
